from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "downloads" / "Enterprise_AI_Governance_Greenfield_Implementation_Blueprint.docx"
NAVY = "102A4C"; GOLD = "A67C2E"; LIGHT = "EEF2F5"; INK = "172D42"; MUTED = "5F6F7E"

stages = [
("01", "Strategy, Vision & Executive Mandate", "Establish why AI governance exists, what outcomes it protects, and which executive has authority.", ["Confirm enterprise AI ambition, risk appetite, regulated obligations and strategic value cases.", "Define scope across predictive AI, generative AI, agentic systems, embedded vendor AI and intelligent automation.", "Approve charter, funding, executive sponsor, governing authority and escalation path."], ["AI governance charter", "AI strategy and principles", "Executive decision record", "Stakeholder map"], "Mandate approved; accountable executive, scope, funding and decision authority are explicit."),
("02", "Current-State Discovery & Maturity Baseline", "Build an evidence-backed view of AI usage, exposure, capabilities and control gaps.", ["Discover sanctioned and unsanctioned AI systems, models, agents, vendors, data flows and automations.", "Assess governance, risk, security, privacy, architecture, data, model operations and adoption maturity.", "Prioritize gaps by business impact, regulatory exposure, autonomy and recoverability."], ["AI inventory", "Maturity assessment", "Gap and dependency register", "Prioritized remediation backlog"], "Inventory owners validate coverage; material unknowns and immediate containment actions are recorded."),
("03", "Governance Operating Model", "Create a four-tier model that connects enterprise authority to delivery controls.", ["Tier 1: Board and executive oversight for appetite, accountability and material risk.", "Tier 2: AI Steering Committee or Centre of Excellence for standards, portfolio decisions and exceptions.", "Tier 3: Domain governance for business ownership, risk acceptance and operational performance.", "Tier 4: Delivery teams for lifecycle evidence, testing, monitoring and incident response."], ["Operating model", "Committee terms of reference", "Decision calendar", "Escalation model"], "Every recurring governance decision has a forum, quorum, cadence, inputs, owner and record."),
("04", "Accountability, RACI & Decision Rights", "Make accountability unambiguous at each lifecycle gate and during incidents.", ["Assign accountable business, model, data, technology, risk, security, privacy and vendor owners.", "Map approval, challenge, consultation, notification and execution rights to lifecycle decisions.", "Define conflict resolution, delegated authority, exception approval and emergency authority."], ["Enterprise RACI", "Decision-rights matrix", "Role descriptions", "Escalation map"], "No material decision has multiple accountable owners or an undefined risk-acceptance authority."),
("05", "Policy, Standards & Control Framework", "Translate principles and obligations into testable controls with evidence requirements.", ["Set acceptable-use, prohibited-use, transparency, human-oversight and third-party requirements.", "Define control objectives across fairness, explainability, privacy, security, data, reliability and monitoring.", "Map controls to NIST AI RMF, ISO/IEC 42001, COBIT, Three Lines, DAMA, model risk and IT service management."], ["AI governance policy", "Standards pack", "Control catalogue", "Framework mapping"], "Each control has an owner, trigger, frequency, evidence, test method, exception route and review date."),
("06", "Lifecycle Governance & Operating Processes", "Embed governance from intake through retirement so control occurs before exposure.", ["Use eight gates: intake, assess, design, develop, validate, deploy, monitor and retire.", "Apply proportional evidence based on risk tier, autonomy, affected stakeholders and reversibility.", "Record approvals, conditions, exceptions and residual risk at each decision gate."], ["Lifecycle playbook", "Intake and risk forms", "Architecture review", "Model card", "Validation report", "Approval and retirement records"], "A system cannot progress without required evidence and an auditable gate decision."),
("07", "AI Risk Management Framework", "Apply proportional governance to operational, behavioural, data, privacy, security and third-party risk.", ["Define taxonomy, inherent risk, control effectiveness, residual risk and acceptance thresholds.", "Link confidence thresholds, human review and automation authority to business consequences.", "Operate risk, issue, exception, incident and remediation registers with formal closure evidence."], ["Risk taxonomy", "Risk assessment", "AI risk register", "Exception register"], "Risk ratings drive controls and approvals; overdue high-risk actions are escalated."),
("08", "Architecture, Data, Security & Privacy", "Build governance into the enterprise control plane and system architecture.", ["Define identity, access, secrets, logging, data lineage, prompt/model gateways and policy enforcement.", "Classify data and enforce purpose, residency, retention, minimization and privacy requirements.", "Test threat scenarios including prompt injection, leakage, poisoning, unsafe tools and excessive agency."], ["Reference architecture", "Architecture decision records", "Data-flow and lineage maps", "Security and privacy assessment"], "Architecture controls are implemented, tested and traceable to requirements before production approval."),
("09", "Enablement, Adoption & Change", "Turn governance into a usable delivery capability rather than a policy-only exercise.", ["Segment training for executives, product owners, risk functions, architects, engineers and users.", "Publish patterns, templates, office hours, intake guidance and communities of practice.", "Measure adoption, decision cycle time, exception causes and recurring evidence gaps."], ["Change strategy", "Role-based curriculum", "Communications plan", "Adoption dashboard"], "Target roles demonstrate competency and delivery teams can use the governance pathway independently."),
("10", "Monitoring, Assurance & Continuous Improvement", "Prove that systems and controls continue to perform after release.", ["Monitor performance, drift, confidence, fairness, security, privacy, cost and business outcomes.", "Apply first-line control operation, second-line oversight and independent third-line assurance.", "Feed incidents, tests, regulatory change and portfolio insights into standards and roadmap updates."], ["KRI/KPI catalogue", "Control testing plan", "Assurance reports", "Incident playbook", "Improvement backlog"], "Material metrics have thresholds and owners; findings are tracked through verified closure."),
]

def font(run, size=11, bold=False, color=INK, name="Aptos"):
    run.font.name=name; run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"),name); run._element.rPr.rFonts.set(qn("w:hAnsi"),name)
    run.font.size=Pt(size); run.bold=bold; run.font.color.rgb=RGBColor.from_string(color)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); tcPr.append(shd)

def set_cell(cell, text, bold=False, color=INK, fill=None):
    cell.text=""; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); r=p.add_run(text); font(r,9,bold,color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill: shade(cell,fill)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.72); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78); sec.header_distance=Inches(.35); sec.footer_distance=Inches(.35)
styles=doc.styles
normal=styles["Normal"]; normal.font.name="Aptos"; normal.font.size=Pt(10); normal.font.color.rgb=RGBColor.from_string(INK); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.2
for key,size,before,after in [("Heading 1",16,18,10),("Heading 2",13,14,7),("Heading 3",11,10,5)]:
    s=styles[key]; s.font.name="Aptos Display"; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(NAVY); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(header.add_run("RATHEESH TECHNOLOGY LTD.  |  ARCHITECTING INTELLIGENCE"),7,True,MUTED)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(footer.add_run("Enterprise AI Governance Greenfield Implementation Blueprint"),7,False,MUTED)
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(10); font(p.add_run("ENTERPRISE IMPLEMENTATION BLUEPRINT"),8,True,GOLD)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(8); font(p.add_run("Greenfield AI Governance\nImplementation Blueprint"),27,True,NAVY,"Aptos Display")
p=doc.add_paragraph(); font(p.add_run("An enterprise-neutral operating model for moving from executive mandate to governed production, assurance, and continuous improvement."),12,False,MUTED)
t=doc.add_table(rows=2,cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
for i,(a,b) in enumerate([("10","implementation stages"),("4","governance tiers"),("8","lifecycle gates"),("3","lines of assurance")]): set_cell(t.cell(0,i),a,True,NAVY,LIGHT); set_cell(t.cell(1,i),b,False,MUTED,LIGHT)
doc.add_heading("How to use this blueprint",1)
doc.add_paragraph("Use the stages as governed workstreams, not a strictly linear project plan. Strategy and discovery establish the mandate; operating model, policy, risk and architecture convert it into controls; lifecycle, enablement and assurance make those controls repeatable. Tailor depth to risk, but retain accountable ownership and auditable decisions.")
doc.add_heading("Ratheesh Technology implementation model",1)
t=doc.add_table(rows=5,cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
for c,v in enumerate(["Mandate","Mobilize","Control","Prove"]): set_cell(t.cell(0,c),v,True,"FFFFFF",NAVY)
rows=[("Strategy and appetite","Inventory and baseline","Policy and standards","Monitoring and assurance"),("Executive authority","Operating model","Risk and lifecycle gates","Independent challenge"),("Scope and funding","RACI and decision rights","Architecture and security","Improvement backlog"),("Outcome: permission to act","Outcome: capability to govern","Outcome: bounded production","Outcome: sustained trust")]
for r,row in enumerate(rows,1):
    for c,v in enumerate(row): set_cell(t.cell(r,c),v,False,INK,"FFFFFF")
doc.add_heading("Implementation stages",1)
for num,title,purpose,activities,deliverables,exitc in stages:
    doc.add_heading(f"{num}  {title}",2); doc.add_paragraph(purpose)
    doc.add_heading("Core activities",3)
    for item in activities: doc.add_paragraph(item,style="List Bullet")
    table=doc.add_table(rows=3,cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    set_cell(table.cell(0,0),"Deliverables",True,"FFFFFF",NAVY); set_cell(table.cell(0,1),"; ".join(deliverables),False,INK,LIGHT)
    set_cell(table.cell(1,0),"Exit criterion",True,NAVY,LIGHT); set_cell(table.cell(1,1),exitc)
    refs={"01":"AI Governance","02":"AI Governance and Resources","03":"AI Governance","04":"AI Governance","05":"Governance, Compliance & Oversight","06":"AI Architecture and Agentic AI Governance","07":"Operational & Behavioural Risk; Data Security & Privacy","08":"Architecture; Security & Architecture Review","09":"AI Adoption and Automation Framework","10":"Assurance"}
    set_cell(table.cell(2,0),"Portal models",True,NAVY,LIGHT); set_cell(table.cell(2,1),refs[num])
doc.add_heading("Implementation roadmap",1)
t=doc.add_table(rows=5,cols=5); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
hdr=["Wave","Typical timing","Primary outcome","Key stage focus","Executive decision"]
for c,v in enumerate(hdr): set_cell(t.cell(0,c),v,True,"FFFFFF",NAVY)
road=[("1 Foundation","Weeks 0–6","Mandate and visibility","01–02","Approve charter and immediate containment"),("2 Mobilize","Weeks 4–12","Accountable governance capability","03–05","Approve model, policy and risk thresholds"),("3 Operationalize","Weeks 8–20","Repeatable lifecycle and controls","06–09","Authorize controlled production pathways"),("4 Prove & scale","Week 16 onward","Evidence, assurance and improvement","10 + all stages","Accept portfolio residual risk and roadmap")]
for r,row in enumerate(road,1):
    for c,v in enumerate(row): set_cell(t.cell(r,c),v,False,INK,LIGHT if r%2==0 else "FFFFFF")
doc.add_heading("Deliverable register",1)
doc.add_paragraph("The companion Excel toolkit contains implementation planning, governance RACI, AI inventory, risk register, control catalogue, lifecycle gates, evidence register, metrics and framework mapping. This Word blueprint provides the narrative operating model and may be adapted into the governance charter, policy suite, committee terms of reference, lifecycle playbook and assurance plan.")
doc.add_heading("Framework alignment",1)
doc.add_paragraph("Use external frameworks as harmonized sources of requirements rather than parallel checklists: NIST AI RMF for Govern/Map/Measure/Manage; ISO/IEC 42001 for the management system; COBIT for enterprise technology governance; the IIA Three Lines Model for ownership, oversight and assurance; DAMA for data governance; model risk management for validation and change control; enterprise architecture for design authority; and IT service management for release, incident, problem and change practices.")
doc.core_properties.title="Greenfield AI Governance Implementation Blueprint"; doc.core_properties.subject="Enterprise AI governance operating model and implementation guide"; doc.core_properties.author="Ratheesh Technology Ltd."
OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT); print(OUT)
