from pathlib import Path
from backend.privacy.gateway import detect_sensitive

LIFECYCLE=["UPLOADED","SCANNING","PII_REVIEW_REQUIRED","PRIVACY_APPROVED","RAG_SAFE_CREATED","CHUNKED","EMBEDDED","VECTOR_VALIDATED","POLICY_OWNER_APPROVED","ACTIVE"]
def inspect_document(path:Path)->dict:
    suffix=path.suffix.lower()
    if suffix==".pdf":
        from pypdf import PdfReader
        text="\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    elif suffix==".docx":
        from docx import Document
        doc=Document(path); text="\n".join([p.text for p in doc.paragraphs]+[" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
    else: text=path.read_text(encoding="utf-8")
    findings=detect_sensitive(text)
    return {"status":"PII_REVIEW_REQUIRED" if findings else "PRIVACY_APPROVED","findings":findings,"text":text}

def may_activate(findings:list[dict],privacy_approved:bool,policy_owner_approved:bool)->bool:
    unresolved=any(f["severity"] in {"HIGH","CRITICAL"} and not f.get("resolved") for f in findings)
    return not unresolved and privacy_approved and policy_owner_approved

def chunk_markdown(text:str)->list[dict]:
    chunks=[]; heading="Introduction"; body=[]
    for line in text.splitlines()+["# EOF"]:
        if line.startswith("#"):
            if body: chunks.append({"section":heading,"content":"\n".join(body).strip(),"status":"ACTIVE"})
            heading=line.lstrip("# ");body=[]
        else: body.append(line)
    return chunks
